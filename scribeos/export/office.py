"""DOCX and EPUB exporters."""

from __future__ import annotations

import html
import io
import re
import uuid

from ..storage import pmdoc
from ..storage.assets import MIME
from ..storage.project_store import Project
from .common import Chapter, ExportOptions, asset_path

HEX6 = re.compile(r"^#([0-9a-fA-F]{6})$")
HEX3 = re.compile(r"^#([0-9a-fA-F]{3})$")


def _rgb(c: str | None):
    from docx.shared import RGBColor

    if not isinstance(c, str):
        return None
    m = HEX6.match(c)
    if m:
        return RGBColor.from_string(m.group(1).upper())
    m = HEX3.match(c)
    if m:
        return RGBColor.from_string("".join(ch * 2 for ch in m.group(1)).upper())
    return None


def _image_stream(project: Project, src: str) -> io.BytesIO | None:
    p = asset_path(project, src)
    if not p:
        return None
    ext = p.suffix[1:]
    if ext in ("png", "jpg"):
        return io.BytesIO(p.read_bytes())
    if ext in ("gif", "webp", "avif"):
        try:
            from PIL import Image

            with Image.open(p) as im:
                im.seek(0)
                buf = io.BytesIO()
                im.convert("RGBA").save(buf, "PNG")
                buf.seek(0)
                return buf
        except Exception:
            return None
    return None


def _add_hyperlink(paragraph, url: str, text: str, run_cb) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = paragraph.add_run(text)
    run_cb(run)
    run.font.underline = True
    run.font.color.rgb = _rgb("#1d4ed8")
    link.append(run._r)
    paragraph._p.append(link)


def export_docx(project: Project, chapters: list[Chapter], opts: ExportOptions) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = opts.font or ("Courier New" if opts.preset == "manuscript" else "Georgia")
    normal.font.size = Pt(opts.fontSize or 12)
    title = opts.title or project.meta.name
    author = opts.author if opts.author is not None else project.meta.author
    d.core_properties.title = title
    d.core_properties.author = author or ""
    align_map = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}

    if opts.pageNumbers:
        footer_p = d.sections[0].footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_p.add_run()
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    if opts.titlePage:
        for _ in range(8):
            d.add_paragraph()
        t = d.add_paragraph(style="Title")
        t.add_run(title)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if author:
            a = d.add_paragraph(author)
            a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if opts.toc and len(chapters) > 1:
        d.add_heading("Contents", level=1)
        p = d.add_paragraph()
        run = p.add_run()
        for kind, text in (("begin", None), (None, 'TOC \\o "1-1" \\h \\z \\u'), ("separate", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        p.add_run("Right-click and choose 'Update Field' to build the table of contents.")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        p.add_run()._r.append(end)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def style_run(run, marks: dict) -> None:
        if "bold" in marks:
            run.bold = True
        if "italic" in marks:
            run.italic = True
        if "underline" in marks:
            run.underline = True
        if "strike" in marks:
            run.font.strike = True
        if "code" in marks:
            run.font.name = "Consolas"
        if "highlight" in marks:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        ts = marks.get("textStyle") or {}
        col = _rgb(ts.get("color"))
        if col is not None:
            run.font.color.rgb = col
        if ts.get("fontFamily"):
            run.font.name = ts["fontFamily"].split(",")[0].strip(" '\"")

    def inline(p, nodes: list) -> None:
        for n in nodes or []:
            if n.get("type") == "hardBreak":
                p.add_run().add_break()
                continue
            marks = {m["type"]: (m.get("attrs") or {}) for m in n.get("marks") or []}
            text = n.get("text", "")
            if "link" in marks and marks["link"].get("href", "").startswith(("http://", "https://", "mailto:")):
                _add_hyperlink(p, marks["link"]["href"], text, lambda r, m=marks: style_run(r, m))
            else:
                style_run(p.add_run(text), marks)

    def block(n: dict, style: str | None = None, level: int = 0) -> None:
        t = n.get("type")
        a = n.get("attrs") or {}
        if t == "paragraph":
            p = d.add_paragraph(style=style)
            if a.get("textAlign") in align_map:
                p.alignment = align_map[a["textAlign"]]
            if opts.preset in ("manuscript", "paperback") and not style:
                p.paragraph_format.first_line_indent = Inches(0.5 if opts.preset == "manuscript" else 0.3)
            inline(p, n.get("content"))
        elif t == "heading":
            h = d.add_heading(level=min(9, int(a.get("level", 2)) + 1))
            inline(h, n.get("content"))
        elif t == "blockquote":
            for c in n.get("content") or []:
                block(c, style="Quote")
        elif t in ("bulletList", "orderedList"):
            base = "List Bullet" if t == "bulletList" else "List Number"
            lvl_style = base if level == 0 else f"{base} {min(3, level + 1)}"
            for li in n.get("content") or []:
                for j, c in enumerate(li.get("content") or []):
                    if c.get("type") in ("bulletList", "orderedList"):
                        block(c, level=level + 1)
                    else:
                        block(c, style=lvl_style if j == 0 else None)
        elif t == "horizontalRule":
            p = d.add_paragraph(opts.sceneBreak)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "codeBlock":
            p = d.add_paragraph()
            r = p.add_run("".join(c.get("text", "") for c in n.get("content") or []))
            r.font.name = "Consolas"
        elif t == "image" and opts.includeImages:
            stream = _image_stream(project, a.get("src", ""))
            if stream:
                width = Inches(6 * (a.get("width") or 100) / 100)
                try:
                    d.add_picture(stream, width=width)
                    d.paragraphs[-1].alignment = align_map.get(a.get("align", "center"), WD_ALIGN_PARAGRAPH.CENTER)
                except Exception:
                    pass
                if a.get("caption"):
                    cp = d.add_paragraph(style="Caption")
                    cp.add_run(a["caption"])
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, ch in enumerate(chapters, 1):
        if i > 1:
            d.add_section(WD_SECTION.NEW_PAGE)
        heading = f"Chapter {i}: {ch.title}" if opts.chapterNumbers else ch.title
        h = d.add_heading(heading, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for n in ch.doc.get("content") or []:
            block(n)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


EPUB_CSS = """
body{font-family:serif;line-height:1.5}
h1{text-align:center;margin:2em 0 1.5em}
p{margin:0;text-indent:1.2em}h1+p,h2+p,h3+p,hr+p,figure+p,blockquote+p{text-indent:0}
blockquote{margin:1em 1.5em;font-style:italic}
figure{margin:1em 0;text-align:center}img{max-width:100%}figcaption{font-size:.9em}
hr{border:0;margin:1.2em 0;text-align:center}hr:after{content:'* * *'}
"""


def export_epub(project: Project, chapters: list[Chapter], opts: ExportOptions) -> bytes:
    from ebooklib import epub

    title = opts.title or project.meta.name
    author = opts.author if opts.author is not None else project.meta.author
    book = epub.EpubBook()
    book.set_identifier(f"urn:uuid:{uuid.uuid5(uuid.NAMESPACE_URL, 'scribeos:' + project.meta.id)}")
    book.set_title(title)
    book.set_language("en")
    if author:
        book.add_author(author)
    css = epub.EpubItem(uid="style", file_name="style/main.css", media_type="text/css", content=EPUB_CSS.encode())
    book.add_item(css)

    added: dict[str, str] = {}

    def resolve(src: str) -> str | None:
        if not opts.includeImages:
            return None
        if src in added:
            return added[src]
        p = asset_path(project, src)
        if not p or p.suffix[1:] not in ("png", "jpg", "gif", "webp", "svg"):
            return None
        fname = f"images/{p.name}"
        book.add_item(epub.EpubItem(uid=f"img-{p.stem[:16]}", file_name=fname, media_type=MIME[p.suffix[1:]], content=p.read_bytes()))
        added[src] = "../" + fname
        return added[src]

    items = []
    if opts.titlePage:
        tp = epub.EpubHtml(title="Title", file_name="text/title.xhtml", lang="en")
        tp.content = f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{html.escape(title)}</title></head><body><h1>{html.escape(title)}</h1>' + (f'<p style="text-align:center;text-indent:0"><em>{html.escape(author)}</em></p>' if author else "") + "</body></html>"
        tp.add_item(css)
        book.add_item(tp)
        items.append(tp)
    chapter_items = []
    for i, ch in enumerate(chapters, 1):
        heading = f"Chapter {i}: {ch.title}" if opts.chapterNumbers else ch.title
        item = epub.EpubHtml(title=heading, file_name=f"text/chapter{i:03d}.xhtml", lang="en")
        body = pmdoc.doc_to_html(ch.doc, resolve, heading_shift=1)
        item.content = f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{html.escape(heading)}</title></head><body><h1>{html.escape(heading)}</h1>{body}</body></html>'
        item.add_item(css)
        book.add_item(item)
        items.append(item)
        chapter_items.append(item)
    book.toc = chapter_items
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = (["nav"] if opts.toc else []) + items
    buf = io.BytesIO()
    epub.write_epub(buf, book, {})
    return buf.getvalue()
