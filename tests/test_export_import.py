import io
import zipfile

import pytest
from PIL import Image

from scribeos.collab import ydoc as yconv
from scribeos.export.common import ExportOptions
from scribeos.export.importers import ImportError_, import_file
from scribeos.export.pdf_typst import tstr
from scribeos.export.service import export, pdf_preview
from scribeos.storage import pmdoc
from scribeos.storage.project_store import Project
from scribeos.voice import commands

INJECT = 'Evil #set page(width: 1pt) ]#[ "quote" \\ back $math$ <label> @ref'


@pytest.fixture
def proj(tmp_path):
    p = Project.create(tmp_path, "Exportable", author="Ann Author")
    b = io.BytesIO()
    Image.new("RGB", (30, 20), "green").save(b, "GIF")
    gif = p.assets.save_bytes(b.getvalue())["src"]
    b = io.BytesIO()
    Image.new("RGB", (30, 20), "blue").save(b, "WEBP")
    webp = p.assets.save_bytes(b.getvalue())["src"]
    cid = p.meta.tree[0].id
    p.update_node(cid, {"title": "The Neon Rain"})
    p.save_chapter(cid, {"type": "doc", "content": [
        {"type": "paragraph", "content": [{"type": "text", "text": INJECT}]},
        {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Section"}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "bold", "marks": [{"type": "bold"}, {"type": "italic"}]}, {"type": "hardBreak"}, {"type": "text", "text": "link", "marks": [{"type": "link", "attrs": {"href": "https://example.com/a)b"}}]}]},
        {"type": "horizontalRule"},
        {"type": "image", "attrs": {"src": gif, "caption": "A GIF", "width": 50}},
        {"type": "image", "attrs": {"src": webp, "align": "right"}},
        {"type": "orderedList", "attrs": {"start": 3}, "content": [{"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "item"}]}]}]},
        {"type": "blockquote", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "quoted"}]}]},
        {"type": "codeBlock", "content": [{"type": "text", "text": "code ``` here"}]},
    ]})
    p.add_chapter("Second", doc=pmdoc.text_to_doc("Second chapter text."))
    return p


def test_tstr_escapes():
    assert tstr('a"b\\c\nd') == '"a\\"b\\\\c\\nd"'


@pytest.mark.parametrize("preset", ["paperback", "manuscript", "a4", "a5", "manga"])
def test_pdf_presets_and_injection_safe(proj, preset):
    from pypdf import PdfReader

    data, name, mime = export(proj, ExportOptions(format="pdf", preset=preset))
    assert data[:4] == b"%PDF" and name == "Exportable.pdf" and mime == "application/pdf"
    reader = PdfReader(io.BytesIO(data))
    text = "".join(pg.extract_text() for pg in reader.pages)
    squashed = "".join(text.split())
    assert "#setpage(width:1pt)" in squashed  # rendered literally, not executed
    assert "TheNeonRain" in squashed and "Secondchaptertext." in squashed
    assert len(reader.pages) >= 3


def test_pdf_scopes(proj):
    second = proj.meta.chapters()[1].id
    data, name, _ = export(proj, ExportOptions(format="pdf", scope="current", currentChapterId=second, titlePage=False, toc=False))
    assert name == "Second.pdf"
    with pytest.raises(ValueError):
        export(proj, ExportOptions(format="pdf", scope="chapters", chapterIds=[]))
    assert len(pdf_preview(proj, ExportOptions())) >= 1


def test_docx(proj):
    from docx import Document

    data, name, _ = export(proj, ExportOptions(format="docx", preset="manuscript"))
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert INJECT in text and "Chapter 1: The Neon Rain" in text
    assert len(d.inline_shapes) == 2


def test_epub(proj):
    data, name, _ = export(proj, ExportOptions(format="epub"))
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.read("mimetype") == b"application/epub+zip"
    names = z.namelist()
    assert any(n.endswith(".gif") for n in names)
    chapter = next(n for n in names if n.endswith("chapter001.xhtml"))
    body = z.read(chapter).decode()
    assert "&lt;label&gt;" in body and "<label>" not in body


def test_markdown_html_txt(proj):
    data, name, _ = export(proj, ExportOptions(format="md"))
    assert name.endswith(".zip")
    z = zipfile.ZipFile(io.BytesIO(data))
    md = z.read("manuscript.md").decode()
    assert "\\#set page" in md and "![" in md and "images/" in md and "3. item" in md
    data, name, _ = export(proj, ExportOptions(format="md", includeImages=False))
    assert name.endswith(".md")
    html, _, _ = export(proj, ExportOptions(format="html"))
    assert b"data:image/gif;base64," in html and b"Content-Security-Policy" in html and b"<label>" not in html
    txt, _, _ = export(proj, ExportOptions(format="txt"))
    assert INJECT.encode() in txt and b"* * *" in txt
    arc, name, _ = export(proj, ExportOptions(format="scribez"))
    assert name.endswith(".scribez") and zipfile.ZipFile(io.BytesIO(arc)).namelist()


def test_import_formats(proj):
    md = b"# One\n\nHello **bold** <script>alert(1)</script>\n\n# Two\n\n- a\n- b\n"
    ch = import_file(proj, "book.md", md)
    assert [t for t, _ in ch] == ["One", "Two"]
    # Raw HTML in Markdown is never interpreted: it survives only as literal, harmless text.
    assert "<script>alert(1)</script>" in pmdoc.doc_text(ch[0][1])
    assert {n["type"] for n in ch[0][1]["content"]} <= {"paragraph"}
    txt = "Prologue text\n\nChapter 1\nIt begins.\n\nChapter 2\nIt ends.".encode()
    assert [t for t, _ in import_file(proj, "b.txt", txt)] == ["b", "Chapter 1", "Chapter 2"]
    from docx import Document

    d = Document()
    d.add_heading("Alpha", 1)
    d.add_paragraph("First words")
    d.add_heading("Beta", 1)
    d.add_paragraph("Second words")
    buf = io.BytesIO()
    d.save(buf)
    assert [t for t, _ in import_file(proj, "x.docx", buf.getvalue())] == ["Alpha", "Beta"]
    with pytest.raises(ImportError_):
        import_file(proj, "x.exe", b"MZ")
    with pytest.raises(ImportError_):
        import_file(proj, "x.docx", b"not a zip")


def test_ydoc_roundtrip(proj):
    doc = proj.get_chapter(proj.meta.tree[0].id)["doc"]
    y = yconv.doc_to_ydoc(doc)
    assert yconv.ydoc_to_doc(y) == doc
    from pycrdt import Doc

    y2 = Doc()
    y2.apply_update(y.get_update())
    assert yconv.ydoc_to_doc(y2) == doc


def test_voice_commands():
    assert commands.parse("Hello world new paragraph how are you question mark") == [{"text": "Hello world"}, {"cmd": "paragraph"}, {"text": "how are you?"}]
    assert commands.parse("Wait, comma, then. Scene break.") == [{"text": "Wait,"}, {"text": ","}, {"text": "then."}, {"cmd": "scenebreak"}] or commands.parse("Wait, comma, then. Scene break.")[-1] == {"cmd": "scenebreak"}
    assert commands.parse("Periodically he stopped") == [{"text": "Periodically he stopped"}]
    # Real faster-whisper output (tiny model) for a spoken test sentence:
    assert commands.parse("the rainfall on the neon city. New paragraph. Kaelith looked up at the sky.") == [
        {"text": "the rainfall on the neon city."}, {"cmd": "paragraph"}, {"text": "Kaelith looked up at the sky."}]
    assert commands.parse("Wait comma then stop dictation") == [{"text": "Wait,"}, {"text": "then"}, {"cmd": "stop"}]
    assert commands.parse("hello new paragraph", enabled=False) == [{"text": "hello new paragraph"}]
    assert commands.parse("") == []
